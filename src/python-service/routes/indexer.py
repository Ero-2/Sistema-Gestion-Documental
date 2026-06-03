import asyncio
import html as _html_escape
import os
from datetime import datetime, timezone

from fastapi import APIRouter, HTTPException, Query, Request
from fastapi.responses import FileResponse, HTMLResponse

from database import collection
from extractor import MAX_CONTENT_CHARS, STORAGE_ROOT, extract_text, extract_file_metadata, get_file_info, resolve_path
from models import PublicDMSMetadata
from pydantic import BaseModel


def _pptx_to_slides_html(path: str) -> str:
    """Convierte un archivo PPTX en HTML de diapositivas usando python-pptx."""
    from pptx import Presentation
    from pptx.util import Pt
    from pptx.enum.text import PP_ALIGN

    prs = Presentation(path)
    slides_html = []

    for i, slide in enumerate(prs.slides, 1):
        shapes_by_y = sorted(
            (s for s in slide.shapes if s.has_text_frame),
            key=lambda s: (s.top or 0),
        )

        title_text = ""
        body_parts = []

        for shape in shapes_by_y:
            tf = shape.text_frame
            lines = []
            for para in tf.paragraphs:
                line = para.text.strip()
                if not line:
                    continue
                size = 0
                for run in para.runs:
                    if run.font.size:
                        size = max(size, run.font.size.pt)
                bold = any(r.font.bold for r in para.runs if r.font.bold is not None)
                lines.append((line, size, bold))

            if not lines:
                continue

            if not title_text and lines:
                title_text = lines[0][0]
                rest = lines[1:]
            else:
                rest = lines

            for text, size, bold in rest:
                tag = "strong" if bold or size >= 20 else "span"
                body_parts.append(f'<{tag}>{_html_escape.escape(text)}</{tag}>')

        slide_num = f'<span class="slide-num">{i}/{len(prs.slides)}</span>'
        title_html = f'<h2>{_html_escape.escape(title_text)}</h2>' if title_text else ''
        body_html = "<br>".join(body_parts)

        slides_html.append(
            f'<div class="slide">{slide_num}{title_html}'
            f'<div class="slide-body">{body_html}</div></div>'
        )

    if not slides_html:
        return '<p style="color:#9aa7b2;padding:2rem">Sin contenido visible en las diapositivas.</p>'

    return "\n".join(slides_html)


class ObsoleteRequest(BaseModel):
    postgres_id: str

router = APIRouter(prefix="/indexer", tags=["Indexer"])

# ── Helpers para documentos simulados (Faker / seed) ─────────────────────────

def _is_seed(file_url: str) -> bool:
    return (file_url or "").startswith("seed/")


def _simulated_content(code: str, title: str, category: str, department: str) -> str:
    """Genera texto fake searchable para docs seed sin archivo físico."""
    prefix = code.split("-")[0] if "-" in code else "DOC"
    blocks = {
        "POL": ["objetivo y alcance", "política general de la organización",
                "responsabilidades de los involucrados", "compromisos institucionales",
                "revisión y actualización periódica"],
        "PRO": ["descripción del procedimiento", "actividades y pasos del proceso",
                "responsables y roles asignados", "registros y evidencias requeridas",
                "frecuencia y puntos de control"],
        "INS": ["instrucciones paso a paso", "herramientas y materiales necesarios",
                "advertencias y precauciones de seguridad", "verificación del resultado",
                "criterios de aceptación"],
        "FOR": ["campos requeridos del formulario", "instrucciones de llenado",
                "datos de identificación y fecha", "firma y aprobación responsable",
                "distribución y archivo del registro"],
    }
    items = blocks.get(prefix, ["contenido general", "descripción del documento",
                                "propósito y aplicación", "vigencia y control"])
    lines = [
        f"{title}",
        f"Código: {code}  |  Categoría: {category}  |  Departamento: {department}",
        "",
        f"Este documento establece los lineamientos para {title.lower()}.",
        f"Aplica al departamento de {department} y a todas las áreas relacionadas con {category}.",
        f"La dirección de {department} es responsable de su implementación y seguimiento.",
        "",
    ]
    for item in items:
        lines.append(f"- {item.capitalize()}: conforme a los estándares de {category}.")
    lines += [
        "",
        "Documento aprobado y vigente según el proceso de gestión de calidad.",
        f"Versión generada automáticamente en el entorno de pruebas para validación de búsqueda full-text.",
    ]
    return "\n".join(lines)


def _simulated_html(doc: dict) -> str:
    title = doc.get("title", "Documento")
    code  = doc.get("code", "—")
    cat   = doc.get("category_name", "—")
    dept  = doc.get("department_name", "—")
    ver   = doc.get("version", "1.0")
    return f"""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<title>{title}</title>
<style>
  body {{ font-family: system-ui, sans-serif; background: #0f172a; color: #e2e8f0;
         display: flex; align-items: center; justify-content: center; min-height: 100vh; margin: 0; }}
  .card {{ background: #1e293b; border: 1px solid #334155; border-radius: 12px;
           padding: 40px 48px; max-width: 520px; text-align: center; }}
  .icon {{ font-size: 3rem; margin-bottom: 16px; }}
  h1 {{ font-size: 1.1rem; color: #f8fafc; margin-bottom: 8px; }}
  p {{ color: #94a3b8; font-size: 0.88rem; line-height: 1.6; margin-bottom: 6px; }}
  .badge {{ display: inline-block; background: #1c3045; color: #7dd3fc;
            border: 1px solid #1e4060; border-radius: 6px; padding: 4px 12px;
            font-size: 0.75rem; font-family: monospace; margin: 12px 0 0; }}
  .meta {{ margin-top: 20px; text-align: left; background: #0f172a; border-radius: 8px;
           padding: 14px 18px; font-size: 0.8rem; }}
  .meta dt {{ color: #64748b; text-transform: uppercase; font-size: 0.7rem; letter-spacing: .04em; }}
  .meta dd {{ color: #cbd5e1; margin: 2px 0 10px; font-family: monospace; }}
</style>
</head>
<body>
<div class="card">
  <div class="icon">📄</div>
  <h1>Documento en entorno de pruebas</h1>
  <p>Este documento fue generado mediante datos simulados (Faker).<br>
     No existe un archivo físico asociado.</p>
  <p>El documento está correctamente indexado y participa en las búsquedas full-text.</p>
  <span class="badge">Documento simulado · sin archivo real</span>
  <dl class="meta">
    <dt>Código</dt><dd>{code}</dd>
    <dt>Título</dt><dd>{title}</dd>
    <dt>Categoría</dt><dd>{cat}</dd>
    <dt>Departamento</dt><dd>{dept}</dd>
    <dt>Versión</dt><dd>{ver}</dd>
  </dl>
</div>
</body>
</html>"""


# ── Endpoints ─────────────────────────────────────────────────────────────────

@router.post("/upsert")
async def upsert_document(metadata: PublicDMSMetadata):
    try:
        doc_data = metadata.model_dump()
        doc_data["sync_date"] = datetime.now(timezone.utc).isoformat()

        # Never overwrite a real file_url with an empty string from a re-sync
        if not metadata.file_url:
            doc_data.pop("file_url", None)

        if metadata.file_url:
            if _is_seed(metadata.file_url):
                # Doc simulado: genera contenido fake para FTS, no intenta leer archivo
                doc_data["file_name"]               = os.path.basename(metadata.file_url)
                doc_data["extension"]               = "pdf"
                doc_data["mime_type"]               = "application/pdf"
                doc_data["size"]                    = 0
                doc_data["path"]                    = metadata.file_url
                doc_data["document_id"]             = metadata.postgres_id
                doc_data["is_simulated"]            = True
                doc_data["content"]                 = _simulated_content(
                    metadata.code, metadata.title,
                    metadata.category_name, metadata.department_name,
                )
                doc_data["content_extracted"]        = True
                doc_data["content_extraction_error"] = None
                doc_data["metadata"] = {
                    "department":   metadata.department_name,
                    "company_id":   metadata.company_id,
                    "company_name": metadata.company_name,
                    "tags":         [],
                    "version":      metadata.version,
                }
            else:
                file_info = get_file_info(metadata.file_url)
                content, extraction_error = await asyncio.to_thread(extract_text, metadata.file_url)
                file_meta = await asyncio.to_thread(extract_file_metadata, metadata.file_url)
                doc_data.update(file_info)
                doc_data["document_id"]              = metadata.postgres_id
                doc_data["is_simulated"]             = False
                doc_data["file_meta"]                = file_meta
                doc_data["metadata"] = {
                    "department":   metadata.department_name,
                    "company_id":   metadata.company_id,
                    "company_name": metadata.company_name,
                    "tags":         [],
                    "version":      metadata.version,
                }
                doc_data["content"]                  = content[:MAX_CONTENT_CHARS]
                doc_data["content_extracted"]        = extraction_error is None
                doc_data["content_extraction_error"] = extraction_error

        # Si ya existe el doc con versión distinta, guardar la versión anterior en el historial
        update_op: dict = {"$set": doc_data}
        existing = await collection.find_one(
            {"postgres_id": metadata.postgres_id},
            {"_id": 0, "version": 1, "file_url": 1, "file_name": 1, "sync_date": 1},
        )
        if existing and existing.get("version") and existing.get("version") != metadata.version:
            old_entry = {
                "version":      existing.get("version"),
                "file_url":     existing.get("file_url", ""),
                "file_name":    existing.get("file_name", ""),
                "obsoleted_at": datetime.now(timezone.utc).isoformat(),
            }
            update_op["$push"] = {"version_history": {"$each": [old_entry], "$position": 0}}

        result = await collection.update_one(
            {"postgres_id": metadata.postgres_id},
            update_op,
            upsert=True,
        )

        return {
            "status":      "success",
            "postgres_id": metadata.postgres_id,
            "action":      "updated" if result.matched_count else "indexed",
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/obsolete")
async def obsolete_document(body: ObsoleteRequest):
    """Marca documento como inactivo en MongoDB cuando es retirado en .NET."""
    postgres_id = body.postgres_id
    try:
        result = await collection.update_one(
            {"postgres_id": str(postgres_id)},
            {"$set": {"is_active": False, "sync_date": datetime.now(timezone.utc).isoformat()}},
        )
        if result.matched_count == 0:
            return {"status": "not_found", "postgres_id": postgres_id}
        return {"status": "success", "postgres_id": postgres_id, "action": "obsoleted"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/search")
async def search_documents(
    q:     str = Query(..., min_length=2, description="Término de búsqueda"),
    limit: int = Query(default=50, le=500),
):
    try:
        cursor = collection.find(
            {"$text": {"$search": q}},
            {"score": {"$meta": "textScore"}, "postgres_id": 1, "_id": 0},
        ).sort([("score", {"$meta": "textScore"})]).limit(limit)

        results = await cursor.to_list(length=limit)
        ids = [int(r["postgres_id"]) for r in results if r.get("postgres_id")]

        return {"ids": ids}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/file/{name}")
async def get_file(name: str):
    """
    Sirve un archivo indexado desde el volumen compartido.
    Si es un documento simulado (seed/Faker), devuelve una página HTML explicativa.
    """
    doc = await collection.find_one(
        {"file_name": name},
        {"_id": 0, "file_url": 1, "path": 1, "file_name": 1, "mime_type": 1,
         "is_simulated": 1, "title": 1, "code": 1,
         "category_name": 1, "department_name": 1, "version": 1},
    )
    _furl = (doc.get("file_url") or doc.get("path", "")) if doc else ""
    if not doc or not _furl:
        raise HTTPException(status_code=404, detail="Documento no encontrado")

    if doc.get("is_simulated") or _is_seed(_furl):
        return HTMLResponse(content=_simulated_html(doc))

    real = os.path.realpath(resolve_path(_furl))
    root = os.path.realpath(STORAGE_ROOT)
    if real != root and not real.startswith(root + os.sep):
        raise HTTPException(status_code=403, detail="Ruta inválida")
    if not os.path.isfile(real):
        raise HTTPException(status_code=404, detail="Archivo no disponible")

    return FileResponse(
        real,
        media_type=doc.get("mime_type") or "application/octet-stream",
        filename=doc.get("file_name") or name,
    )


@router.get("/download/{name}")
async def download_file(name: str):
    """
    Fuerza descarga del archivo (Content-Disposition: attachment).
    Docs simulados devuelven 422 con mensaje descriptivo.
    """
    doc = await collection.find_one(
        {"file_name": name},
        {"_id": 0, "file_url": 1, "path": 1, "file_name": 1, "mime_type": 1, "is_simulated": 1},
    )
    _furl = (doc.get("file_url") or doc.get("path", "")) if doc else ""
    if not doc or not _furl:
        raise HTTPException(status_code=404, detail="Documento no encontrado")

    if doc.get("is_simulated") or _is_seed(_furl):
        raise HTTPException(
            status_code=422,
            detail="Documento simulado — no existe archivo físico para descargar. "
                   "Fue generado mediante datos de prueba (Faker).",
        )

    real = os.path.realpath(resolve_path(_furl))
    root = os.path.realpath(STORAGE_ROOT)
    if real != root and not real.startswith(root + os.sep):
        raise HTTPException(status_code=403, detail="Ruta inválida")
    if not os.path.isfile(real):
        raise HTTPException(status_code=404, detail="Archivo no disponible")

    return FileResponse(
        real,
        media_type=doc.get("mime_type") or "application/octet-stream",
        filename=doc.get("file_name") or name,
        headers={"Content-Disposition": f'attachment; filename="{doc.get("file_name") or name}"'},
    )


@router.get("/viewer/{name}", response_class=HTMLResponse, include_in_schema=False)
async def viewer_page(name: str, request: Request):
    """
    Visor de documentos en el navegador. Obtiene metadatos de MongoDB,
    embebe el contenido extraído cuando aplica y devuelve HTML que renderiza
    el archivo según su extensión.
    """
    # Prefijo público detrás de Nginx por path (/fastapi).
    base = request.headers.get("x-forwarded-prefix", "").rstrip("/")
    doc = await collection.find_one(
        {"file_name": name},
        {"_id": 0, "title": 1, "code": 1, "extension": 1, "mime_type": 1,
         "is_simulated": 1, "content": 1, "content_extracted": 1,
         "content_extraction_error": 1, "category_name": 1, "department_name": 1,
         "version": 1, "file_url": 1, "path": 1},
    )

    import json as _json

    if not doc:
        raise HTTPException(status_code=404, detail="Documento no encontrado")

    title   = doc.get("title") or name
    code    = doc.get("code") or ""
    ext     = (doc.get("extension") or "").lower().lstrip(".")
    isSim   = doc.get("is_simulated", False)
    content = doc.get("content") or ""
    fileUrl = f"{base}/indexer/file/{name}"
    dlUrl   = f"{base}/indexer/download/{name}"

    simInfo = ""
    if isSim:
        cat  = doc.get("category_name", "—")
        dept = doc.get("department_name", "—")
        ver  = doc.get("version", "1.0")
        simInfo = f"""
        <div class="sim-card">
          <div class="sim-icon">📄</div>
          <h2>Documento en entorno de pruebas</h2>
          <p>Este documento fue generado con datos simulados (Faker).<br>
             No existe un archivo físico asociado.</p>
          <p>El documento está indexado y participa en búsquedas full-text.</p>
          <span class="sim-badge">Documento simulado · sin archivo real</span>
          <dl class="sim-meta">
            <dt>Código</dt><dd>{code}</dd>
            <dt>Título</dt><dd>{title}</dd>
            <dt>Categoría</dt><dd>{cat}</dd>
            <dt>Departamento</dt><dd>{dept}</dd>
            <dt>Versión</dt><dd>{ver}</dd>
          </dl>
        </div>"""

    html = f"""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title}</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=IBM+Plex+Mono:wght@400;500;600&family=IBM+Plex+Sans:wght@400;500;600&display=swap" rel="stylesheet">
<style>
:root {{
  --bg: #0e1217; --surface-1: #141a21; --surface-2: #19212a;
  --tx-1: #d7dee5; --tx-2: #9aa7b2; --tx-3: #6b7682; --tx-4: #4a535c;
  --accent: #4fd1c5; --danger: #e0727a;
  --bd-1: rgba(200,215,225,0.08); --bd-2: rgba(200,215,225,0.14); --bd-3: rgba(200,215,225,0.22);
  --mono: 'IBM Plex Mono', ui-monospace, monospace;
  --sans: 'IBM Plex Sans', system-ui, sans-serif;
}}
* {{ box-sizing: border-box; margin: 0; padding: 0; }}
html, body {{ height: 100%; }}
body {{ background: var(--bg); color: var(--tx-1); font-family: var(--sans); font-size: 14px; -webkit-font-smoothing: antialiased; }}
/* ── Topbar ── */
.topbar {{
  display: flex; align-items: center; gap: 10px;
  background: var(--surface-1); border-bottom: 1px solid var(--bd-2);
  padding: 10px 20px; height: 52px; position: sticky; top: 0; z-index: 10;
}}
.btn-back {{
  font-family: var(--mono); font-size: 12px; color: var(--tx-2);
  background: transparent; border: 1px solid var(--bd-2);
  border-radius: 4px; padding: 6px 12px; cursor: pointer;
  transition: border-color .12s, color .12s; flex-shrink: 0;
}}
.btn-back:hover {{ border-color: var(--bd-3); color: var(--tx-1); }}
.topbar-title {{ font-family: var(--mono); font-size: 13px; font-weight: 500; color: var(--tx-1); white-space: nowrap; overflow: hidden; text-overflow: ellipsis; flex: 1; min-width: 0; }}
.ext-badge {{
  font-family: var(--mono); font-size: 11px; font-weight: 600;
  background: #1e3a5f; color: #7dd3fc;
  border-radius: 4px; padding: 3px 8px; flex-shrink: 0;
}}
.btn-dl {{
  font-family: var(--mono); font-size: 12px; color: var(--tx-2);
  background: transparent; border: 1px solid var(--bd-2);
  border-radius: 4px; padding: 6px 14px; cursor: pointer;
  transition: border-color .12s, color .12s; flex-shrink: 0; text-decoration: none;
  display: inline-flex; align-items: center; gap: 6px;
}}
.btn-dl:hover {{ border-color: var(--accent); color: var(--accent); }}
/* ── Viewer area ── */
#viewer {{ height: calc(100vh - 52px); overflow: auto; }}
iframe {{ width: 100%; height: 100%; border: none; }}
#image-wrap {{ display: flex; justify-content: center; align-items: flex-start; padding: 32px; }}
#image-wrap img {{ max-width: 100%; border-radius: 6px; box-shadow: 0 4px 24px rgba(0,0,0,.5); }}
#docx-wrap {{ padding: 2rem; max-width: 860px; margin: 0 auto; background: #fff; color: #111; min-height: 100%; font-family: serif; }}
#sheet-tabs {{ background: var(--surface-1); border-bottom: 1px solid var(--bd-2); padding: 6px 14px; display: flex; gap: 6px; flex-wrap: wrap; }}
#sheet-tabs button {{ background: var(--surface-2); border: 1px solid var(--bd-2); color: var(--tx-2); padding: 4px 12px; border-radius: 4px; font-size: 12px; font-family: var(--mono); cursor: pointer; }}
#sheet-tabs button.active {{ background: var(--accent); border-color: var(--accent); color: #0e1217; }}
#sheet-wrap {{ padding: 1rem; overflow: auto; height: calc(100vh - 52px - 40px); }}
#sheet-wrap table {{ border-collapse: collapse; font-size: .82rem; font-family: var(--mono); }}
#sheet-wrap th {{ background: var(--surface-2); color: var(--tx-2); padding: 6px 10px; border: 1px solid var(--bd-2); white-space: nowrap; }}
#sheet-wrap td {{ border: 1px solid var(--bd-1); padding: 5px 10px; color: var(--tx-1); }}
#text-wrap {{ padding: 24px 28px; }}
/* ── PPTX slides ── */
#pptx-wrap {{ padding: 24px 32px; max-width: 900px; margin: 0 auto; }}
.slide {{
  background: var(--surface-1); border: 1px solid var(--bd-2);
  border-radius: 8px; padding: 28px 32px; margin-bottom: 20px;
  position: relative;
}}
.slide-num {{ position: absolute; top: 10px; right: 14px; font-family: var(--mono); font-size: 11px; color: var(--tx-4); }}
.slide h2 {{ font-size: 1.15rem; font-weight: 600; color: var(--tx-1); margin-bottom: 14px; padding-bottom: 10px; border-bottom: 1px solid var(--bd-1); }}
.slide-body {{ font-size: .9rem; color: var(--tx-2); line-height: 1.7; }}
.slide-body strong {{ color: var(--tx-1); font-weight: 600; display: block; margin-top: 6px; }}
#text-wrap pre {{ background: var(--surface-1); border: 1px solid var(--bd-1); border-radius: 6px; padding: 16px; font-family: var(--mono); font-size: 13px; color: #a8e6cf; white-space: pre-wrap; word-break: break-word; line-height: 1.6; }}
.preview-note {{
  background: var(--surface-2); border: 1px solid var(--bd-2);
  border-radius: 6px; padding: 10px 16px; margin-bottom: 16px;
  font-family: var(--mono); font-size: 12px; color: var(--tx-3);
  display: flex; align-items: center; gap: 10px;
}}
.preview-note a {{ color: var(--accent); text-decoration: none; margin-left: auto; }}
.preview-note a:hover {{ text-decoration: underline; }}
/* ── Unsupported / sim ── */
.center-page {{ display: flex; align-items: center; justify-content: center; height: calc(100vh - 52px); }}
.unsup-card {{
  background: var(--surface-1); border: 1px solid var(--bd-2);
  border-radius: 10px; padding: 36px 44px; max-width: 420px; text-align: center;
}}
.unsup-icon {{ font-size: 2.8rem; margin-bottom: 14px; }}
.unsup-card h3 {{ color: var(--tx-1); font-size: 1rem; margin-bottom: 10px; }}
.unsup-card p {{ color: var(--tx-3); font-size: 0.85rem; line-height: 1.6; margin-bottom: 16px; }}
.unsup-card .dl-btn {{
  display: inline-flex; align-items: center; gap: 8px;
  font-family: var(--mono); font-size: 13px; color: var(--tx-1);
  background: var(--surface-2); border: 1px solid var(--bd-2);
  border-radius: 6px; padding: 9px 20px; text-decoration: none;
  transition: border-color .12s, color .12s;
}}
.unsup-card .dl-btn:hover {{ border-color: var(--accent); color: var(--accent); }}
/* simulado */
.sim-card {{
  background: var(--surface-1); border: 1px solid var(--bd-2);
  border-radius: 12px; padding: 40px 48px; max-width: 500px; text-align: center;
}}
.sim-icon {{ font-size: 3rem; margin-bottom: 14px; }}
.sim-card h2 {{ font-size: 1.05rem; color: var(--tx-1); margin-bottom: 10px; }}
.sim-card p {{ color: var(--tx-3); font-size: .875rem; line-height: 1.6; margin-bottom: 8px; }}
.sim-badge {{
  display: inline-block; font-family: var(--mono); font-size: 10px;
  color: #7dd3fc; background: #1c3045; border: 1px solid #1e4060;
  border-radius: 4px; padding: 4px 12px; margin: 12px 0;
}}
.sim-meta {{ margin-top: 20px; text-align: left; background: var(--bg); border-radius: 6px; padding: 14px 18px; font-size: .8rem; }}
.sim-meta dt {{ color: var(--tx-4); text-transform: uppercase; font-size: .7rem; letter-spacing: .04em; }}
.sim-meta dd {{ color: var(--tx-2); margin: 2px 0 10px; font-family: var(--mono); }}
/* spinner */
.spin-wrap {{ display: flex; justify-content: center; align-items: center; padding: 64px; }}
.spinner {{ width: 28px; height: 28px; border: 3px solid var(--bd-2); border-top-color: var(--accent); border-radius: 50%; animation: spin .8s linear infinite; }}
@keyframes spin {{ to {{ transform: rotate(360deg); }} }}
</style>
</head>
<body>

<div class="topbar">
  <button class="btn-back" onclick="history.length > 1 ? history.back() : (location.href = '{base}/search')">&#8592; volver</button>
  <span class="topbar-title">{title}</span>
  {'<span class="ext-badge">' + ext.upper() + '</span>' if ext else ''}
  {'<a href="' + dlUrl + '" class="btn-dl">&#8595; descargar</a>' if not isSim else ''}
</div>

<div id="viewer">
"""

    # ── Render según extensión ────────────────────────────────────────────────
    IMAGE_EXTS = {"png", "jpg", "jpeg", "gif", "webp", "bmp", "svg", "avif", "ico"}
    TEXT_EXTS  = {"txt", "csv", "log", "json", "xml", "yaml", "yml", "ini", "md", "tsv", "conf"}

    if isSim:
        html += f'<div class="center-page">{simInfo}</div>'

    elif ext == "pdf":
        html += f'<iframe src="{fileUrl}" title="{title}"></iframe>'

    elif ext in IMAGE_EXTS:
        html += f'<div id="image-wrap"><img src="{fileUrl}" alt="{title}"></div>'

    elif ext in {"html", "htm"}:
        html += f'<iframe src="{fileUrl}" sandbox="allow-same-origin" title="{title}"></iframe>'

    elif ext == "docx":
        html += f"""<div id="docx-wrap"><div class="spin-wrap"><div class="spinner"></div></div></div>
<script src="https://cdn.jsdelivr.net/npm/mammoth@1.8.0/mammoth.browser.min.js"></script>
<script>
fetch({_json.dumps(fileUrl)})
  .then(r => r.arrayBuffer())
  .then(buf => mammoth.convertToHtml({{ arrayBuffer: buf }}))
  .then(r => {{ document.getElementById('docx-wrap').innerHTML = r.value || '<p style="color:#888">Sin contenido visible.</p>'; }})
  .catch(e => {{ document.getElementById('docx-wrap').innerHTML = '<p style="color:#e0727a">Error al procesar DOCX: ' + e.message + '</p>'; }});
</script>"""

    elif ext in {"xlsx", "xls"}:
        html += f"""<div id="sheet-tabs"></div>
<div id="sheet-wrap"><div class="spin-wrap"><div class="spinner"></div></div></div>
<script src="https://cdn.jsdelivr.net/npm/xlsx@0.18.5/dist/xlsx.full.min.js"></script>
<script>
let _wb;
fetch({_json.dumps(fileUrl)})
  .then(r => r.arrayBuffer())
  .then(buf => {{
    _wb = XLSX.read(buf, {{ type: 'array' }});
    const tabs = document.getElementById('sheet-tabs');
    _wb.SheetNames.forEach((n, i) => {{
      const b = document.createElement('button');
      b.textContent = n;
      if (i === 0) b.classList.add('active');
      b.onclick = () => {{
        document.querySelectorAll('#sheet-tabs button').forEach(x => x.classList.remove('active'));
        b.classList.add('active');
        renderSheet(n);
      }};
      tabs.appendChild(b);
    }});
    renderSheet(_wb.SheetNames[0]);
  }})
  .catch(e => {{ document.getElementById('sheet-wrap').innerHTML = '<p style="color:#e0727a;padding:1rem">Error al procesar Excel: ' + e.message + '</p>'; }});
function renderSheet(n) {{
  document.getElementById('sheet-wrap').innerHTML = XLSX.utils.sheet_to_html(_wb.Sheets[n], {{ editable: false }});
}}
</script>"""

    elif ext in {"pptx", "ppt"}:
        _pptx_url = doc.get("file_url") or doc.get("path", "")
        real = os.path.realpath(resolve_path(_pptx_url))
        root = os.path.realpath(STORAGE_ROOT)
        if os.path.isfile(real) and (real == root or real.startswith(root + os.sep)):
            try:
                slides_html = await asyncio.to_thread(_pptx_to_slides_html, real)
                html += f'<div id="pptx-wrap">{slides_html}</div>'
            except Exception as e:
                import html as _h
                html += f'<div id="text-wrap"><div class="preview-note">Error al renderizar PPTX: {_h.escape(str(e))} <a href="{dlUrl}">&#8595; Descargar</a></div></div>'
        elif content and doc.get("content_extracted"):
            import html as _h
            safe = _h.escape(content[:80000])
            html += f'<div id="text-wrap"><div class="preview-note">Vista previa de texto (.{ext}) <a href="{dlUrl}">&#8595; Descargar original</a></div><pre>{safe}</pre></div>'
        else:
            html += f"""<div class="center-page"><div class="unsup-card">
  <div class="unsup-icon">📊</div>
  <h3>Archivo no disponible</h3>
  <p>El archivo <strong>.{ext}</strong> no se encontró en el servidor.</p>
  <a href="{dlUrl}" class="dl-btn">&#8595; Descargar {name}</a>
</div></div>"""

    elif ext in TEXT_EXTS:
        # Texto disponible directamente desde MongoDB si fue extraído
        if content and doc.get("content_extracted"):
            import html as _html_mod
            safe = _html_mod.escape(content[:80000])
            html += f'<div id="text-wrap"><pre>{safe}</pre></div>'
        else:
            html += f"""<div id="text-wrap">
<div class="preview-note">Cargando texto… <a href="{dlUrl}">&#8595; Descargar original</a></div>
<pre id="txt">…</pre>
</div>
<script>
fetch({_json.dumps(fileUrl)})
  .then(r => r.text())
  .then(t => {{ document.getElementById('txt').textContent = t; document.querySelector('.preview-note').style.display='none'; }})
  .catch(e => {{ document.getElementById('txt').textContent = 'Error: ' + e.message; }});
</script>"""

    else:
        # Fallback: texto extraído por el indexador o prompt de descarga
        if content and doc.get("content_extracted"):
            import html as _html_mod
            safe = _html_mod.escape(content[:80000])
            html += f"""<div id="text-wrap">
<div class="preview-note">Vista previa de texto (.{ext} no se renderiza en el navegador).
  <a href="{dlUrl}">&#8595; Descargar original</a>
</div>
<pre>{safe}</pre>
</div>"""
        else:
            html += f"""<div class="center-page">
<div class="unsup-card">
  <div class="unsup-icon">📎</div>
  <h3>Vista previa no disponible</h3>
  <p>El formato <strong>.{ext}</strong> no puede renderizarse en el navegador.<br>
     Descarga el archivo para abrirlo con la aplicación correspondiente.</p>
  <a href="{dlUrl}" class="dl-btn">&#8595; Descargar {name}</a>
</div>
</div>"""

    html += "\n</div>\n</body>\n</html>"
    return HTMLResponse(content=html)


@router.get("/content/{name}")
async def get_content(name: str):
    """
    Texto extraído de un documento indexado. Para docs simulados devuelve el contenido
    generado automáticamente.
    """
    doc = await collection.find_one(
        {"file_name": name},
        {"_id": 0, "title": 1, "extension": 1, "content": 1,
         "content_extracted": 1, "content_extraction_error": 1, "is_simulated": 1},
    )
    if not doc:
        raise HTTPException(status_code=404, detail="Documento no encontrado")

    return {
        "title":       doc.get("title"),
        "extension":   doc.get("extension"),
        "extracted":   bool(doc.get("content_extracted")),
        "error":       doc.get("content_extraction_error"),
        "is_simulated": bool(doc.get("is_simulated")),
        "content":     doc.get("content") or "",
    }
